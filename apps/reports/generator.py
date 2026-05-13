import csv
import io
from decimal import Decimal
from django.db.models import Sum
from apps.transactions.models import Transaction


class ReportGenerator:
    def __init__(self, user, date_from, date_to):
        self.user = user
        self.date_from = date_from
        self.date_to = date_to
        self._qs = None

    @property
    def transactions(self):
        if self._qs is None:
            self._qs = (
                Transaction.objects
                .filter(user=self.user, date__gte=self.date_from, date__lte=self.date_to)
                .select_related('category')
                .only('date', 'transaction_type', 'description', 'amount', 'tags', 'category__name')
                .order_by('date')
            )
        return self._qs

    def generate_csv(self):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['Date', 'Type', 'Category', 'Description', 'Amount', 'Tags'])
        # iterator() avoids loading all rows into memory at once
        for txn in self.transactions.iterator(chunk_size=500):
            writer.writerow([
                txn.date, txn.transaction_type,
                txn.category.name if txn.category else '',
                txn.description, txn.amount, txn.tags,
            ])
        totals = self.transactions.values('transaction_type').annotate(total=Sum('amount'))
        writer.writerow([])
        writer.writerow(['Summary'])
        for row in totals:
            writer.writerow([row['transaction_type'].title(), '', '', 'Total', row['total'], ''])
        return output.getvalue()

    def generate_excel(self):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook(write_only=True)  # write_only=True for memory efficiency
        ws = wb.create_sheet('Transactions')

        header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        headers = ['Date', 'Type', 'Category', 'Description', 'Amount', 'Tags']

        header_row = []
        for header in headers:
            from openpyxl.cell.cell import WriteOnlyCell
            cell = WriteOnlyCell(ws, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            header_row.append(cell)
        ws.append(header_row)

        for txn in self.transactions.iterator(chunk_size=500):
            ws.append([
                str(txn.date),
                txn.transaction_type.title(),
                txn.category.name if txn.category else '',
                txn.description,
                float(txn.amount),
                txn.tags,
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def get_summary(self):
        # Single aggregation query for both income and expense
        rows = self.transactions.values('transaction_type').annotate(total=Sum('amount'))
        totals = {r['transaction_type']: r['total'] or Decimal('0') for r in rows}
        income = totals.get('income', Decimal('0'))
        expense = totals.get('expense', Decimal('0'))
        by_category = list(
            self.transactions
            .values('category__name', 'transaction_type')
            .annotate(total=Sum('amount'))
            .order_by('-total')
        )
        return {'income': income, 'expense': expense, 'net': income - expense, 'by_category': by_category}
